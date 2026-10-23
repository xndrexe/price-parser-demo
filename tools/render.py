"""Рисует снимки таблиц для README из настоящих данных последнего запуска.

.venv/bin/python tools/render.py   # после двух запусков python -m pricewatch; HTML → build/screens/
Нужны Excel и хотя бы одно изменение в базе. «Товары» — одна категория (Travel),
«Изменения» — свежие сверху, без столбца ссылок; уведомление снято в настоящем Telegram.
"""

from __future__ import annotations

import sys
from html import escape
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from pricewatch import config, export, store  # noqa: E402

OUT = ROOT / "build" / "screens"

CSS = """
* { box-sizing: border-box; margin: 0; }
body { font: 15px/1.38 Inter, 'Noto Color Emoji', sans-serif; background: #fff; }
table { border-collapse: collapse; font: 12px/1.3 Inter, sans-serif; margin: 12px; }
td, th { border: 1px solid #d4d4d4; padding: 5px 8px; white-space: nowrap; max-width: 330px;
  overflow: hidden; text-overflow: ellipsis; }
th { background: #2f5597; color: #fff; font-weight: 600; text-align: left; }
tr:nth-child(odd) td { background: #f3f6fb; }
td.num { text-align: right; }
.sheet { padding: 0 0 4px; background: #fff; display: inline-block; }
.sheet h1 { font: 600 14px Inter, sans-serif; padding: 10px 12px 0; color: #333; }
"""


def page(body: str) -> str:
    return f"<!doctype html><meta charset=utf-8><style>{CSS}</style>{body}"


def render_sheet(path: Path, sheet: str, rows: int, skip_cols: tuple[int, ...] = (), only: str = "") -> str:
    """only — показать строки одной категории (второй столбец), чтобы снимок был нагляднее."""
    ws = load_workbook(path)[sheet]
    data = list(ws.iter_rows(values_only=True))
    head, body = data[0], [r for r in data[1:] if not only or r[1] == only]
    out = []
    for i, row in enumerate([head, *body[:rows]]):
        tag = "th" if i == 0 else "td"
        cells = []
        for j, v in enumerate(row):
            if j in skip_cols:
                continue
            num = isinstance(v, int | float)
            text = f"{v:.2f}" if isinstance(v, float) else "" if v is None else str(v)
            cls = ' class="num"' if num and tag == "td" else ""
            cells.append(f"<{tag}{cls}>{escape(text)}</{tag}>")
        out.append("<tr>" + "".join(cells) + "</tr>")
    shown = f"категория {only}" if only else f"первые {min(rows, len(body))}"
    return page(
        f'<div class="sheet"><h1>Лист «{sheet}» — {len(data) - 1} строк, на снимке {shown}</h1>'
        f"<table>{''.join(out)}</table></div>"
    )


def main() -> None:
    cfg = config.load()
    conn = store.connect(cfg.db_path)
    rows = store.recent_changes(conn, limit=1)
    conn.close()
    if not rows:
        sys.exit("нет изменений в базе: запусти парсер дважды")
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "2-tablitsa.html").write_text(
        render_sheet(cfg.excel_path, export.PRODUCTS_SHEET, 14, skip_cols=(6,), only="Travel"), "utf-8"
    )
    (OUT / "3-izmeneniya.html").write_text(
        render_sheet(cfg.excel_path, export.CHANGES_SHEET, 10, skip_cols=(5,)), "utf-8"
    )
    for f in sorted(OUT.glob("*.html")):
        print(f)


if __name__ == "__main__":
    main()
