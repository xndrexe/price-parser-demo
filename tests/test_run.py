"""Сквозной запуск без сети: обход → база → Excel → Telegram, два прогона подряд."""

from dataclasses import replace

import pytest
from conftest import make_book, make_list
from openpyxl import load_workbook

from pricewatch import __main__ as app
from pricewatch import config

ROOT = "https://shop.test/catalogue/"


class FakeHTTP:
    """Подменяет requests.Session: страницы сайта по GET, Telegram по POST."""

    def __init__(self, pages):
        self.pages = pages
        self.sent = []

    def get(self, url, timeout):
        page = self.pages.get(url)

        class R:
            status_code = 200 if page is not None else 404
            encoding = "utf-8"
            text = page or ""

            def raise_for_status(self):
                if page is None:
                    raise OSError(f"404 {url}")

        return R()

    def post(self, url, json, timeout):
        self.sent.append(json["text"])

        class R:
            def raise_for_status(self):
                pass

        return R()


def shop(price_a=10.0, stock_b=4):
    return {
        ROOT + "page-1.html": make_list(["a/index.html", "b/index.html"], None),
        ROOT + "a/index.html": make_book("upc-a", "Книга А", price_a, 7),
        ROOT + "b/index.html": make_book("upc-b", "Книга Б", 20.0, stock_b),
    }


@pytest.fixture
def cfg(tmp_path):
    return replace(
        config.load(),
        start_url=ROOT + "page-1.html",
        max_pages=0,
        workers=2,
        delay=0,
        db_path=tmp_path / "db.sqlite",
        excel_path=tmp_path / "out" / "products.xlsx",
        tg_token="1:T",
        tg_chat_id="42",
        gsheet_id="",
        google_credentials="",
    )


def test_two_runs_end_to_end(cfg):
    http = FakeHTTP(shop())
    r1 = app.run_once(cfg, http)
    assert (r1.products, r1.changes, r1.complete) == (2, 0, True)
    assert http.sent == []

    wb = load_workbook(cfg.excel_path)
    rows = list(wb["Товары"].values)
    assert rows[0][0] == "Название"
    assert sorted(r[0] for r in rows[1:]) == ["Книга А", "Книга Б"]

    http.pages = shop(price_a=8.0, stock_b=0)
    r2 = app.run_once(cfg, http)
    assert r2.changes == 2
    assert len(http.sent) == 1
    assert "£10.00 → £8.00" in http.sent[0] and "закончился" in http.sent[0]
    changes = list(load_workbook(cfg.excel_path)["Изменения"].values)
    assert {r[1] for r in changes[1:]} == {"Цена", "Наличие"}
    price = next(r for r in changes[1:] if r[1] == "Цена")
    assert (price[3], price[4]) == (10.0, 8.0)  # числом, а не текстом — работают фильтры


def test_dead_site_keeps_old_table(cfg):
    http = FakeHTTP(shop())
    app.run_once(cfg, http)
    before = cfg.excel_path.read_bytes()
    http.pages = {}
    with pytest.raises(RuntimeError):
        app.run_once(cfg, http)
    assert cfg.excel_path.read_bytes() == before


def test_schedule_survives_failed_run(monkeypatch):
    """С --every упавший проход пишется в журнал, расписание не останавливается."""
    runs = []

    def flaky(cfg):
        runs.append(1)
        if len(runs) == 1:
            raise RuntimeError("site down")
        return app.RunReport(1, 0, 0, True)

    class Stop(Exception):
        pass

    def sleep(sec):
        assert sec == 60 * 60
        if len(runs) == 2:
            raise Stop

    monkeypatch.setattr(app, "run_once", flaky)
    monkeypatch.setattr(app.time, "sleep", sleep)
    with pytest.raises(Stop):
        app.main(["--every", "60"])
    assert len(runs) == 2


def test_single_run_failure_is_loud(monkeypatch):
    monkeypatch.setattr(app, "run_once", lambda cfg: (_ for _ in ()).throw(RuntimeError("site down")))
    with pytest.raises(RuntimeError):
        app.main([])


def test_pages_flag_overrides_config(monkeypatch, cfg):
    seen = []
    monkeypatch.setattr(app.config, "load", lambda: cfg)
    monkeypatch.setattr(app, "run_once", lambda c: seen.append(c.max_pages) or app.RunReport(0, 0, 0, True))
    app.main(["--pages", "5"])
    app.main([])
    assert seen == [5, 0]


def test_no_telegram_without_token(cfg):
    http = FakeHTTP(shop())
    quiet = replace(cfg, tg_token="")
    app.run_once(quiet, http)
    http.pages = shop(price_a=8.0)
    assert app.run_once(quiet, http).changes == 1
    assert http.sent == []


def test_partial_run_reports_errors(cfg):
    pages = shop()
    del pages[ROOT + "b/index.html"]
    r = app.run_once(cfg, FakeHTTP(pages))
    assert (r.products, r.errors, r.complete) == (1, 1, False)


def test_telegram_failure_keeps_run_and_changes(cfg, monkeypatch, caplog):
    """Telegram лёг — запуск не падает, изменения остаются на листе «Изменения»."""
    http = FakeHTTP(shop())
    app.run_once(cfg, http)

    def down(session, token, chat_id, text):
        raise OSError(f"401 for url: https://api.telegram.org/bot{token}/sendMessage")

    monkeypatch.setattr(app.notify, "send_telegram", down)
    http.pages = shop(price_a=8.0)
    with caplog.at_level("WARNING"):
        assert app.run_once(cfg, http).changes == 1
    assert "telegram send failed" in caplog.text and cfg.tg_token not in caplog.text
    changes = list(load_workbook(cfg.excel_path)["Изменения"].values)
    assert [r[1] for r in changes[1:]] == ["Цена"]


@pytest.mark.parametrize("argv", [["--every", "-5"], ["--pages", "-1"]])
def test_negative_flags_rejected(argv, monkeypatch):
    monkeypatch.setattr(app, "run_once", lambda c: pytest.fail("не должен запускаться"))
    with pytest.raises(SystemExit) as e:
        app.main(argv)
    assert e.value.code == 2
