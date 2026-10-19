"""Выгрузка в Excel и подготовка строк — без сети."""

import pytest
from openpyxl import load_workbook

from pricewatch import export, store
from pricewatch.parse import Product


@pytest.mark.parametrize(
    "raw, value",
    [
        (None, ""),
        ("22", 22),
        ("0", 0),
        ("10.00", 10.0),
        ("8.50", 8.5),
        ("-3", -3.0),
        ("нет", "нет"),
        ("", ""),
    ],
)
def test_num(raw, value):
    got = export._num(raw)
    assert got == value and type(got) is type(value)


def test_change_rows_translate_kind_and_keep_unknown():
    rows = export.change_rows(
        [
            {"ts": "2026-10-15T09:00:00", "kind": "gone", "title": "К", "old": None, "new": None, "url": "u"},
            {"ts": "2026-10-15T09:00:00", "kind": "rating", "title": "К", "old": "3", "new": "4", "url": "u"},
        ]
    )
    assert rows[0] == ["2026-10-15 09:00:00", "Пропал с сайта", "К", "", "", "u"]
    assert rows[1][1:5] == ["rating", "К", 3, 4]


def test_to_excel_layout_and_no_tmp_left(tmp_path):
    conn = store.connect(tmp_path / "db.sqlite")
    store.save_snapshot(conn, [Product("a", "Книга", 1234.5, 3, 4, "Поэзия", "https://shop.test/a")], True)
    path = tmp_path / "deep" / "dir" / "products.xlsx"  # папки создаются сами
    assert export.to_excel(store.current(conn), store.recent_changes(conn), path) == path
    assert [p.name for p in path.parent.iterdir()] == ["products.xlsx"]

    wb = load_workbook(path)
    assert wb.sheetnames == ["Товары", "Изменения"]
    goods = wb["Товары"]
    assert [c.value for c in goods[1]] == export.PRODUCT_HEADERS
    assert goods.freeze_panes == "A2" and goods.auto_filter.ref == "A1:H2"
    price = goods["C2"]
    assert price.value == 1234.5 and price.number_format == "#,##0.00"
    assert "T" not in goods["H2"].value  # дата читаемая, без «T»
    hist = wb["Изменения"]
    assert [c.value for c in hist[1]] == export.CHANGE_HEADERS and hist.max_row == 1


def test_to_excel_overwrites_previous_file(tmp_path):
    path = tmp_path / "products.xlsx"
    conn = store.connect(tmp_path / "db.sqlite")
    store.save_snapshot(conn, [Product("a", "Старая", 1, 1, 1, "К", "u")], True)
    export.to_excel(store.current(conn), [], path)
    store.save_snapshot(conn, [Product("a", "Новая", 1, 1, 1, "К", "u")], True)
    export.to_excel(store.current(conn), [], path)
    assert load_workbook(path)["Товары"]["A2"].value == "Новая"


def test_control_chars_do_not_break_excel(tmp_path):
    # управляющий символ в названии раньше ронял всю выгрузку
    conn = store.connect(tmp_path / "db.sqlite")
    for price in (10.0, 9.0):
        p = Product("a", "Книга\x0bодин", price, 3, 4, "По\x00эзия", "https://s.test/a")
        store.save_snapshot(conn, [p], True)
    path = export.to_excel(store.current(conn), store.recent_changes(conn), tmp_path / "x.xlsx")
    wb = load_workbook(path)
    assert wb["Товары"]["A2"].value == "Книгаодин" and wb["Товары"]["B2"].value == "Поэзия"
    assert wb["Изменения"]["C2"].value == "Книгаодин"


@pytest.mark.parametrize(
    "raw, value",
    [
        ("036000291452", "'036000291452"),  # артикул из цифр не теряет ведущий ноль
        ("=IMPORTXML(1)", "'=IMPORTXML(1)"),
        ("+7", "'+7"),
        ("Книга", "Книга"),
        ("", ""),
        (12.5, 12.5),
        (7, 7),
    ],
)
def test_gsheet_text(raw, value):
    assert export._gsheet_text(raw) == value
