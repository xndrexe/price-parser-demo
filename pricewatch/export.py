"""Выгрузка: файл Excel и, если настроено, Google Таблица."""

from __future__ import annotations

import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PRODUCT_HEADERS = [
    "Название",
    "Категория",
    "Цена, £",
    "В наличии, шт",
    "Рейтинг",
    "Артикул (UPC)",
    "Ссылка",
    "Обновлено",
]
CHANGE_HEADERS = ["Когда", "Что", "Товар", "Было", "Стало", "Ссылка"]
KIND_RU = {
    "price": "Цена",
    "stock": "Наличие",
    "new": "Новый товар",
    "gone": "Пропал с сайта",
    "back": "Снова на сайте",
}


def product_rows(products: list[sqlite3.Row]) -> list[list]:
    return [
        [
            p["title"],
            p["category"],
            p["price"],
            p["stock"],
            p["rating"],
            p["upc"],
            p["url"],
            p["last_seen"].replace("T", " "),
        ]
        for p in products
    ]


def _num(v: str | None) -> float | int | str:
    """Было/стало хранятся текстом; в таблицу — числом, чтобы работали фильтры и формулы."""
    if v is None:
        return ""
    try:
        return int(v) if v.isdigit() else float(v)
    except ValueError:
        return v


def change_rows(changes: list[sqlite3.Row]) -> list[list]:
    return [
        [
            c["ts"].replace("T", " "),
            KIND_RU.get(c["kind"], c["kind"]),
            c["title"],
            _num(c["old"]),
            _num(c["new"]),
            c["url"],
        ]
        for c in changes
    ]


FORMULA_START = ("=", "+", "-", "@")


def _gsheet_text(v):
    """Текст с сайта, похожий на формулу, — в Google Таблицу как текст (апостроф Таблица не показывает).

    Иначе при USER_ENTERED название вида «=IMPORTXML(...)» выполнится у всех, кто откроет таблицу.
    """
    if isinstance(v, str) and (v.startswith(FORMULA_START) or v.isdigit()):
        # строка из одних цифр (артикул «036000291452») иначе станет числом без ведущего нуля
        return "'" + v
    return v


def _sheet(ws, headers: list[str], rows: list[list], widths: list[int]) -> None:
    ws.append(headers)
    for row in rows:
        # управляющие символы из вёрстки (\x0b и т. п.) Excel не принимает: без чистки падает
        # вся выгрузка, а изменения уже в базе — уведомление о них потерялось бы
        ws.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in row])
        # openpyxl считает формулой любую строку с «=» в начале; данные с сайта — всегда текст
        for cell in ws[ws.max_row]:
            if cell.data_type == "f":
                cell.data_type = "s"
    head_fill = PatternFill("solid", fgColor="2F5597")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def to_excel(products: list[sqlite3.Row], changes: list[sqlite3.Row], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"
    _sheet(ws, PRODUCT_HEADERS, product_rows(products), [48, 18, 10, 14, 9, 20, 40, 19])
    for (cell,) in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        cell.number_format = "#,##0.00"
    _sheet(wb.create_sheet("Изменения"), CHANGE_HEADERS, change_rows(changes), [19, 16, 48, 10, 10, 40])
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)  # таблица не окажется полузаписанной, если её как раз открыли
    return path


def to_gsheet(
    products: list[sqlite3.Row], changes: list[sqlite3.Row], sheet_id: str, credentials: str
) -> None:
    """Перезаливает листы «Товары» и «Изменения» в Google Таблицу.

    credentials — путь к ключу сервисного аккаунта Google; таблицу нужно
    расшарить на его почту с правом редактирования.
    """
    import gspread  # нужен только при выгрузке в Google

    book = gspread.service_account(filename=credentials).open_by_key(sheet_id)
    for title, headers, rows in (
        ("Товары", PRODUCT_HEADERS, product_rows(products)),
        ("Изменения", CHANGE_HEADERS, change_rows(changes)),
    ):
        try:
            ws = book.worksheet(title)
        except gspread.WorksheetNotFound:
            ws = book.add_worksheet(title, rows=max(len(rows) + 1, 10), cols=len(headers))
        ws.clear()
        safe = [[_gsheet_text(v) for v in row] for row in rows]
        ws.update([headers, *safe], value_input_option="USER_ENTERED")
        ws.freeze(rows=1)
